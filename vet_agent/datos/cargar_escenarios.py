"""Carga los escenarios de Investigación (Excel) y los guarda como JSON.

El Excel "plantilla_escenario-cargable.xlsx" trae una hoja "Escenarios" donde
cada fila es un escenario. Las columnas azules son la Parte A (lo único que
recibe el sistema) y las naranjas la Parte B (la respuesta esperada, que el
sistema nunca debe ver). Este módulo convierte cada fila en un JSON con las dos
partes separadas y lo guarda en corpus/escenarios/IAAP-XX.json.

Reglas que sigue (ver hoja "Instrucciones" del Excel):
- Los nombres de columna usan puntos para indicar anidación:
  "parte_a_entrada.clima.temperatura_media_c" -> {"parte_a_entrada": {"clima":
  {"temperatura_media_c": ...}}}.
- Las columnas con varios elementos vienen separadas por ";" y se convierten en
  lista (signos clínicos, factores de riesgo, documentos esperados, etc.).
- resultados_laboratorio_1 y resultados_laboratorio_2 se juntan en una lista de
  objetos {"prueba", "resultado", "fecha"}.
- Las celdas vacías pasan a null y la última columna (tipo_caso) se ignora.

Uso:
    python -m vet_agent.datos.cargar_escenarios <ruta_del_excel>
"""

import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

HOJA = "Escenarios"
DIR_SALIDA = Path(__file__).resolve().parents[2] / "corpus" / "escenarios"

# Columnas cuyos valores vienen separados por ";" y se convierten en lista.
CAMPOS_LISTA = {
    "parte_a_entrada.signos_clinicos",
    "parte_b_referencia.factores_riesgo",
    "parte_b_referencia.hipotesis_alternativas",
    "parte_b_referencia.informacion_faltante",
    "parte_b_referencia.medidas_inmediatas",
    "parte_b_referencia.protocolos_aplicables",
    "parte_b_referencia.documentos_esperados",
    "parte_b_referencia.acciones_que_no_debe_recomendar",
}

# Columnas numéricas y su tipo de conversión (n_animales es entero, el resto float).
CAMPOS_NUMERICOS = {
    "parte_a_entrada.n_animales": int,
    "parte_a_entrada.mortalidad_pct": float,
    "parte_a_entrada.morbilidad_pct": float,
    "parte_a_entrada.clima.temperatura_media_c": float,
    "parte_a_entrada.clima.humedad_relativa_pct": float,
    "parte_a_entrada.clima.precipitacion_mm": float,
}

# Columnas de resultados de laboratorio: se juntan en una lista de objetos.
PREFIJO_RESULTADOS = "parte_a_entrada.resultados_laboratorio_"

# Última columna del Excel: la usa Investigación (E2) para su revisión; se ignora.
COLUMNA_IGNORADA = "tipo_caso"


def _normalizar_id(valor) -> str:
    """Limpia el id del escenario. El ejemplo viene como 'EJEMPLO IAAP-00' y
    nos quedamos con 'IAAP-00'."""
    texto = str(valor).strip() if valor is not None else ""
    prefijo = "EJEMPLO "
    if texto.upper().startswith(prefijo):
        texto = texto[len(prefijo):]
    return texto


def _a_numero(celda, tipo) -> int | float | None:
    """Convierte una celda a número; acepta texto con '%' o espacios."""
    if celda is None:
        return None
    if isinstance(celda, str):
        limpio = celda.replace("%", "").replace(",", ".").strip()
        if not limpio:
            return None
        return tipo(limpio)
    return tipo(celda)


def _texto_a_lista(celda) -> list | None:
    """'a; b; c' -> ['a', 'b', 'c']; celda vacía -> None."""
    if celda is None:
        return None
    if isinstance(celda, str):
        elementos = [e.strip() for e in celda.split(";")]
        elementos = [e for e in elementos if e]
        return elementos or None
    return [celda]


def _celda_a_texto(celda) -> str | None:
    """Celda -> texto, o None si está vacía. Fechas de Excel -> AAAA-MM-DD."""
    if celda is None:
        return None
    if isinstance(celda, datetime):
        return celda.strftime("%Y-%m-%d")
    if isinstance(celda, date):
        return celda.isoformat()
    texto = str(celda).strip()
    return texto or None


def _valor_de_celda(campo: str, celda):
    """Convierte el valor de una celda según el tipo de su columna."""
    if campo in CAMPOS_NUMERICOS:
        return _a_numero(celda, CAMPOS_NUMERICOS[campo])
    if campo in CAMPOS_LISTA:
        return _texto_a_lista(celda)
    return _celda_a_texto(celda)


def _poner(destino: dict, ruta: str, valor) -> None:
    """Pone `valor` en destino siguiendo la ruta punteada: 'a.b.c' -> destino['a']['b']['c']."""
    partes = ruta.split(".")
    nodo = destino
    for parte in partes[:-1]:
        nodo = nodo.setdefault(parte, {})
    nodo[partes[-1]] = valor


def convertir_fila(cabecera, fila) -> dict:
    """Convierte una fila del Excel (lista de valores) en el JSON del escenario."""
    escenario: dict = {}
    resultados: dict[int, dict] = {}

    for nombre, celda in zip(cabecera, fila):
        if nombre is None:
            continue
        nombre = str(nombre).strip()
        if not nombre or nombre.lower().startswith(COLUMNA_IGNORADA.lower()):
            continue

        if nombre.startswith(PREFIJO_RESULTADOS):
            # 'parte_a_entrada.resultados_laboratorio_1.prueba' -> resultados[1]['prueba']
            resto = nombre[len(PREFIJO_RESULTADOS):]  # '1.prueba'
            num, _, campo = resto.partition(".")
            resultados.setdefault(int(num), {})[campo] = _celda_a_texto(celda)
            continue

        _poner(escenario, nombre, _valor_de_celda(nombre, celda))

    # juntamos los resultados de laboratorio en una lista de objetos
    lista_resultados = []
    for num in sorted(resultados):
        res = resultados[num]
        if all(v is None for v in res.values()):
            continue  # las tres celdas vacías: no hay resultado
        lista_resultados.append({
            "prueba": res.get("prueba"),
            "resultado": res.get("resultado"),
            "fecha": res.get("fecha"),
        })
    _poner(escenario, "parte_a_entrada.resultados_laboratorio",
           lista_resultados if lista_resultados else None)

    if "id_escenario" in escenario:
        escenario["id_escenario"] = _normalizar_id(escenario["id_escenario"])

    return escenario


def escenarios_de_hoja(hoja: Worksheet) -> list[dict]:
    """Lee todas las filas de la hoja (cabecera + filas) y devuelve un dict por escenario."""
    filas = list(hoja.iter_rows(values_only=True))
    if not filas:
        return []
    cabecera = filas[0]
    escenarios = []
    for fila in filas[1:]:
        vacia = all(c is None or (isinstance(c, str) and not c.strip()) for c in fila)
        if vacia:
            continue
        escenarios.append(convertir_fila(cabecera, fila))
    return escenarios


def leer_escenarios(ruta_excel) -> list[dict]:
    """Abre el Excel y devuelve la lista de escenarios de la hoja 'Escenarios'."""
    libro = load_workbook(ruta_excel, data_only=True)
    return escenarios_de_hoja(libro[HOJA])


def guardar_escenarios(ruta_excel, carpeta_salida=None) -> list[dict]:
    """Lee el Excel y guarda un JSON por escenario en `carpeta_salida` (por defecto
    corpus/escenarios). Devuelve la lista de escenarios leídos."""
    escenarios = leer_escenarios(ruta_excel)
    carpeta = Path(carpeta_salida) if carpeta_salida else DIR_SALIDA
    carpeta.mkdir(parents=True, exist_ok=True)
    for escenario in escenarios:
        id_esc = escenario.get("id_escenario")
        if not id_esc:
            continue
        ruta = carpeta / f"{id_esc}.json"
        ruta.write_text(json.dumps(escenario, ensure_ascii=False, indent=2) + "\n",
                        encoding="utf-8")
    return escenarios


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("Falta la ruta del Excel. Uso: python -m vet_agent.datos.cargar_escenarios <ruta_del_excel>")
    escenarios = guardar_escenarios(sys.argv[1])
    print(f"Se generaron {len(escenarios)} escenarios en {DIR_SALIDA}")
    for escenario in escenarios:
        print(f"  - {escenario.get('id_escenario')}.json")
